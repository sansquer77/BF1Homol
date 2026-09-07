import json
import re
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests" / "fixtures" / "backups" / "v3_excel"
MANIFEST = json.loads((FIXTURES / "manifest.json").read_text(encoding="utf-8"))
SCHEMA = json.loads((FIXTURES / "schema.json").read_text(encoding="utf-8"))
EXPECTED_PASSWORD = "$2b$04$r.mdeGukB1b4FFqUP9IquOMh/Eda1S/4EDrvx19EYUgJnRp4zQ08q"


def _rows(table: str):
    workbook = openpyxl.load_workbook(FIXTURES / f"{table}.xlsx", read_only=True, data_only=False)
    assert workbook.sheetnames == ["data"]
    return list(workbook["data"].iter_rows(values_only=True))


def test_excel_fixture_set_matches_versioned_contract():
    assert len(MANIFEST["tables"]) == 21
    assert set(path.stem for path in FIXTURES.glob("*.xlsx")) == set(MANIFEST["tables"])
    for table, contract in MANIFEST["tables"].items():
        rows = _rows(table)
        assert list(rows[0]) == contract["columns"]
        assert len(rows) - 1 == contract["rows"]


def test_reconstructed_schema_accepts_every_excel_contract_column():
    assert SCHEMA["contract"] == "v3.x-reconstructed"
    assert set(SCHEMA["tables"]) == set(MANIFEST["tables"])
    for table, contract in MANIFEST["tables"].items():
        schema_columns = {column["column_name"] for column in SCHEMA["tables"][table]["columns"]}
        assert set(contract["columns"]) <= schema_columns


def test_excel_fixtures_redact_authentication_and_personal_data():
    users = _rows("usuarios")
    headers = list(users[0])
    records = [dict(zip(headers, row)) for row in users[1:]]
    assert records
    assert all(str(row["email"]).endswith("@example.invalid") for row in records)
    assert all(str(row["nome"]).startswith("Participante ") for row in records)
    assert all(row["senha_hash"] == EXPECTED_PASSWORD for row in records)

    for table in ("access_logs", "login_attempts"):
        rows = _rows(table)
        headers = list(rows[0])
        for row in rows[1:]:
            record = dict(zip(headers, row))
            if record.get("email"):
                assert str(record["email"]).endswith("@example.invalid")
            if record.get("ip_address"):
                assert re.fullmatch(r"192\.0\.2\.(?:[1-9]|[1-9]\d|1\d\d|2[0-4]\d|25[0-4])", str(record["ip_address"]))


def test_excel_fixture_foreign_key_values_are_in_parent_exports():
    def values(table: str, column: str) -> set:
        rows = _rows(table)
        index = list(rows[0]).index(column)
        return {row[index] for row in rows[1:] if row[index] is not None}

    users = values("usuarios", "id")
    races = values("provas", "id")
    rules = values("regras", "id")
    assert values("apostas", "usuario_id") <= users
    assert values("apostas", "prova_id") <= races
    assert values("posicoes_participantes", "usuario_id") <= users
    assert values("posicoes_participantes", "prova_id") <= races
    assert values("temporadas_regras", "regra_id") <= rules
